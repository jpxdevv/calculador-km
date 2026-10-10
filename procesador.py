"""
Lógica de cálculo de kilometraje, rendimiento y faltantes de diésel.

Reglas:
- Las cargas del mismo tracto en la misma fecha se juntan en una sola línea:
  se suman litros y total, y los tickets quedan en una sola celda.
- KM FINAL     = odómetro de la línea (el mayor si hubo varios tickets ese día).
- KM INICIAL   = KM FINAL de la última línea anterior del mismo tracto con
                 FOLIO LIQ. > 0.
- KM RECORRIDOS = KM FINAL - KM INICIAL.
- Rendimiento meta sale de la hoja catálogo (columna REND META SB); si el tracto
  no está, se usa META_DEFAULT.

Uso desde terminal:
    python procesador.py "archivo.xlsx"
"""
import sys
from pathlib import Path

import pandas as pd

META_DEFAULT = 2.5

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO",
         "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

COLUMNAS_REQUERIDAS = ["TICKET", "FECHA", "TRACTO", "LTS", "FOLIO LIQ.", "ODOMETRO"]


class ErrorDeDatos(Exception):
    pass


def corregir_texto(valor):
    # Repara acentos dañados por codificación (ej. "SALDAÃ‘A" -> "SALDAÑA").
    if isinstance(valor, str) and "Ã" in valor:
        for codificacion in ("cp1252", "latin-1"):
            try:
                return valor.encode(codificacion).decode("utf-8")
            except (UnicodeEncodeError, UnicodeDecodeError):
                pass
    return valor


def unir_unicos(serie):
    valores = [str(v) for v in pd.unique(serie.dropna())]
    return ", ".join(valores) if valores else None


def preparar_bd(df):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    # La columna del odómetro puede venir como "ODOMETRO" o "kms".
    df = df.rename(columns={c: "ODOMETRO" for c in df.columns
                            if c.upper() in ("ODOMETRO", "KMS", "KM")})
    faltan = [c for c in COLUMNAS_REQUERIDAS if c not in df.columns]
    if faltan:
        raise ErrorDeDatos("Al Excel le faltan estas columnas: " + ", ".join(faltan))

    for col in df.select_dtypes(include=["object", "str"]):
        df[col] = df[col].map(corregir_texto)

    df["FECHA"] = pd.to_datetime(df["FECHA"], errors="coerce", dayfirst=True)
    for col in ("ODOMETRO", "LTS", "TOTAL", "FOLIO LIQ.", "TRACTO"):
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df


def leer_metas(hojas, nombre_bd):
    for nombre, hoja in hojas.items():
        if nombre == nombre_bd:
            continue
        hoja = hoja.rename(columns=lambda c: str(c).strip())
        if "TRACTO" in hoja.columns and "REND META SB" in hoja.columns:
            metas = hoja[["TRACTO", "REND META SB"]].copy()
            metas["TRACTO"] = pd.to_numeric(metas["TRACTO"], errors="coerce")
            metas["REND META SB"] = pd.to_numeric(metas["REND META SB"], errors="coerce")
            metas = metas.dropna().drop_duplicates("TRACTO")
            return dict(zip(metas["TRACTO"], metas["REND META SB"]))
    return {}


def agrupar_por_dia(df):
    """Una línea por tracto y fecha."""
    df = df.copy()
    df["_orden"] = range(len(df))
    agg = {
        "_orden": "min",
        "TICKET": unir_unicos,
        "LTS": "sum",
        "ODOMETRO": "max",
        "FOLIO LIQ.": "max",
        "FOLIOS LIQ.": unir_unicos,
        "NO. TICKETS": "size",
    }
    df["FOLIOS LIQ."] = df["FOLIO LIQ."].map(lambda v: f"{v:.0f}" if pd.notna(v) else None)
    df["NO. TICKETS"] = 1
    if "TOTAL" in df.columns:
        agg["TOTAL"] = "sum"
    for col in ("FLOTA", "GASOLINERA", "RAZON", "REMOLQUE", "OPERADOR", "NOM OPERADOR",
                "PRE-LIQ", "FACTURA", "TIPO CARGA", "FOLIO CUADRE", "CAPTURO"):
        if col in df.columns:
            agg[col] = unir_unicos

    lineas = df.groupby(["TRACTO", "FECHA"], as_index=False, dropna=False).agg(agg)
    lineas = lineas.rename(columns={"TICKET": "TICKETS"})
    if "TOTAL" in lineas.columns:
        lineas["PRECIO PROM."] = (lineas["TOTAL"] / lineas["LTS"]).round(2)
    return lineas


def calcular_km(lineas, metas):
    lineas = lineas.sort_values(["TRACTO", "FECHA", "_orden"]).copy()
    lineas["MES"] = lineas["FECHA"].dt.month.map(lambda m: MESES[int(m) - 1] if pd.notna(m) else None)
    lineas["KM FINAL"] = lineas["ODOMETRO"]

    con_folio = (lineas["FOLIO LIQ."] > 0) & lineas["ODOMETRO"].notna()
    ultimo = lineas["ODOMETRO"].where(con_folio)
    # Último km liquidado hacia abajo por tracto, desplazado una línea (el *anterior*).
    lineas["KM INICIAL"] = ultimo.groupby(lineas["TRACTO"]).transform(lambda s: s.ffill().shift())
    lineas["KM RECORRIDOS"] = lineas["KM FINAL"] - lineas["KM INICIAL"]

    lineas["REND. META"] = lineas["TRACTO"].map(metas).fillna(META_DEFAULT)
    lineas["META POR DEFECTO"] = ~lineas["TRACTO"].isin(list(metas))
    km_validos = lineas["KM RECORRIDOS"].where(lineas["KM RECORRIDOS"] > 0)
    lineas["RENDIMIENTO"] = (km_validos / lineas["LTS"].where(lineas["LTS"] > 0)).round(2)
    lineas["LTS META"] = (km_validos / lineas["REND. META"]).round(2)
    lineas["FALTANTE LTS"] = (lineas["LTS"] - lineas["LTS META"]).round(2)
    if "PRECIO PROM." in lineas.columns:
        lineas["FALTANTE $"] = (lineas["FALTANTE LTS"] * lineas["PRECIO PROM."]).round(2)

    obs = pd.Series("", index=lineas.index)
    obs[lineas["KM INICIAL"].isna()] = "SIN KM INICIAL (no hay línea previa liquidada)"
    obs[lineas["KM RECORRIDOS"] < 0] = "REVISAR: KM MENOR AL ANTERIOR"
    obs[lineas["KM RECORRIDOS"] == 0] = "REVISAR: 0 KM RECORRIDOS"
    lineas["OBSERVACIONES"] = obs
    return lineas.sort_values(["FECHA", "_orden"])


def resumen_por_folio(lineas):
    con_folio = lineas[lineas["FOLIO LIQ."] > 0].sort_values(["TRACTO", "FECHA", "_orden"])
    if con_folio.empty:
        return pd.DataFrame()
    validas = con_folio["KM RECORRIDOS"] > 0
    con_folio = con_folio.assign(
        _km=con_folio["KM RECORRIDOS"].where(validas),
        _lts=con_folio["LTS"].where(validas),
        _lts_meta=con_folio["LTS META"].where(validas),
        _falt_pesos=con_folio["FALTANTE $"].where(validas) if "FALTANTE $" in con_folio else None,
    )
    g = con_folio.groupby("FOLIO LIQ.")
    res = pd.DataFrame({
        "TRACTO": g["TRACTO"].first(),
        "OPERADOR": g["OPERADOR"].first() if "OPERADOR" in con_folio else None,
        "NOM OPERADOR": g["NOM OPERADOR"].first() if "NOM OPERADOR" in con_folio else None,
        "FECHA INICIAL": g["FECHA"].min(),
        "FECHA FINAL": g["FECHA"].max(),
        "KM INICIAL": g["KM INICIAL"].min(),
        "KM FINAL": g["KM FINAL"].max(),
        "KM RECORRIDOS": g["_km"].sum(min_count=1),
        "LTS CARGADOS": g["LTS"].sum(),
        "LTS CON KM": g["_lts"].sum(min_count=1),
        "REND. META": g["REND. META"].first(),
        "LTS META": g["_lts_meta"].sum(min_count=1),
        "NO. TICKETS": g["NO. TICKETS"].sum(),
        "TICKETS": g["TICKETS"].agg(lambda s: ", ".join(s.dropna())),
    }).reset_index()
    res["RENDIMIENTO"] = (res["KM RECORRIDOS"] / res["LTS CON KM"]).round(2)
    res["FALTANTE LTS"] = (res["LTS CON KM"] - res["LTS META"]).round(2)
    if "FALTANTE $" in con_folio:
        res["FALTANTE $"] = g["_falt_pesos"].sum(min_count=1).values.round(2)
    res["FOLIO LIQ."] = res["FOLIO LIQ."].astype("int64")
    orden = ["FOLIO LIQ.", "TRACTO", "OPERADOR", "NOM OPERADOR", "FECHA INICIAL", "FECHA FINAL",
             "KM INICIAL", "KM FINAL", "KM RECORRIDOS", "LTS CARGADOS", "LTS CON KM", "RENDIMIENTO",
             "REND. META", "LTS META", "FALTANTE LTS", "FALTANTE $", "NO. TICKETS", "TICKETS"]
    return res[[c for c in orden if c in res.columns]].sort_values("FOLIO LIQ.")


def procesar(ruta_entrada):
    """Lee el Excel y regresa un dict con las tablas de resultado y un resumen."""
    try:
        hojas = pd.read_excel(ruta_entrada, sheet_name=None)
    except Exception as e:
        raise ErrorDeDatos(f"No se pudo leer el archivo: {e}") from e
    nombre_bd = next(iter(hojas))
    bd = preparar_bd(hojas[nombre_bd])
    metas = leer_metas(hojas, nombre_bd)

    sin_km = bd[bd["ODOMETRO"].isna()]
    lineas = calcular_km(agrupar_por_dia(bd[bd["ODOMETRO"].notna()]), metas)
    resumen = resumen_por_folio(lineas)

    orden = ["MES", "FECHA", "TRACTO", "OPERADOR", "NOM OPERADOR", "NO. TICKETS", "TICKETS",
             "FOLIOS LIQ.", "KM INICIAL", "KM FINAL", "KM RECORRIDOS", "LTS", "PRECIO PROM.", "TOTAL",
             "RENDIMIENTO", "REND. META", "META POR DEFECTO", "LTS META", "FALTANTE LTS", "FALTANTE $",
             "OBSERVACIONES", "FLOTA", "GASOLINERA", "RAZON", "TIPO CARGA", "PRE-LIQ", "FACTURA",
             "FOLIO CUADRE", "CAPTURO"]
    detalle = lineas[[c for c in orden if c in lineas.columns]]
    revisar = detalle[detalle["OBSERVACIONES"].str.startswith("REVISAR")]

    totales = {
        "tickets": len(bd),
        "lineas": len(detalle),
        "folios": len(resumen),
        "km": float(resumen["KM RECORRIDOS"].sum()) if len(resumen) else 0.0,
        "faltante_lts": float(resumen.loc[resumen["FALTANTE LTS"] > 0, "FALTANTE LTS"].sum()) if len(resumen) else 0.0,
        "faltante_pesos": float(resumen.loc[resumen.get("FALTANTE $", pd.Series(dtype=float)) > 0, "FALTANTE $"].sum())
        if len(resumen) and "FALTANTE $" in resumen else 0.0,
        "revisar": len(revisar),
        "sin_km": len(sin_km),
        "tractos_sin_meta": sorted(int(t) for t in detalle.loc[detalle["META POR DEFECTO"], "TRACTO"].unique()),
    }
    return {
        "Resumen por folio": resumen,
        "Detalle por día": detalle,
        "Revisar": revisar,
        "Sin km (remolque)": sin_km,
        "_totales": totales,
    }


def guardar(resultado, ruta_salida):
    from openpyxl.styles import Font, PatternFill

    rojo = PatternFill("solid", fgColor="F8D7DA")
    with pd.ExcelWriter(ruta_salida, engine="openpyxl", datetime_format="DD/MM/YYYY") as xl:
        for nombre, tabla in resultado.items():
            if nombre.startswith("_") or tabla is None or tabla.empty:
                continue
            tabla.to_excel(xl, sheet_name=nombre[:31], index=False)
            ws = xl.sheets[nombre[:31]]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for celda in ws[1]:
                celda.font = Font(bold=True)
            for col in ws.columns:
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(ancho + 2, 8), 45)
            # Marca en rojo los faltantes positivos.
            encabezados = [c.value for c in ws[1]]
            if "FALTANTE LTS" in encabezados:
                idx = encabezados.index("FALTANTE LTS") + 1
                for fila in ws.iter_rows(min_row=2):
                    v = fila[idx - 1].value
                    if isinstance(v, (int, float)) and v > 0:
                        for c in fila:
                            c.fill = rojo


def ruta_salida_para(ruta_entrada):
    p = Path(ruta_entrada)
    return p.with_name(f"{p.stem} - RESULTADO.xlsx")


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    entrada = sys.argv[1]
    resultado = procesar(entrada)
    salida = ruta_salida_para(entrada)
    guardar(resultado, salida)
    t = resultado["_totales"]
    print(f"Listo: {salida}")
    print(f"  Tickets: {t['tickets']}  |  Líneas por día: {t['lineas']}  |  Folios: {t['folios']}")
    print(f"  Faltante: {t['faltante_lts']:,.2f} L  /  ${t['faltante_pesos']:,.2f}")
    print(f"  Para revisar: {t['revisar']}  |  Tractos con meta por defecto: {len(t['tractos_sin_meta'])}")


if __name__ == "__main__":
    main()
